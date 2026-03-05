from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.models import TradeRecord

BUY_FONT = Font(color="00C00000")
SELL_FONT = Font(color="00004BBA")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="00EFEFEF")


def _fmt(value: int) -> str:
    return f"{value:,}"


def _summary_totals(trades: Iterable[TradeRecord]) -> Tuple[int, int]:
    buy_total = 0
    sell_total = 0
    for trade in trades:
        if trade.side == "buy":
            buy_total += trade.total_amount
        elif trade.side == "sell":
            sell_total += trade.total_amount
    return buy_total, sell_total


def generate_excel(trades: Iterable[TradeRecord], trade_date: str, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "매매 요약"
    ws_detail = wb.create_sheet("체결 상세")

    summary_headers = ["종목명", "구분", "총수량", "가중평균가", "총 체결금액", "매매 근거"]
    detail_headers = ["주문번호", "종목명", "구분", "체결시간", "체결수량", "체결가", "체결금액", "시장"]

    ws_summary.append(summary_headers)
    ws_detail.append(detail_headers)

    for col in "ABCDEF":
        ws_summary[f"{col}1"].fill = HEADER_FILL
    for col in "ABCDEFGH":
        ws_detail[f"{col}1"].fill = HEADER_FILL

    trade_list = list(trades)

    for row_idx, trade in enumerate(trade_list, start=2):
        side_label = "매수" if trade.side == "buy" else "매도"
        ws_summary.append(
            [
                trade.stock_name,
                side_label,
                trade.total_qty,
                trade.avg_price,
                trade.total_amount,
                trade.reason,
            ]
        )

        font = BUY_FONT if trade.side == "buy" else SELL_FONT
        for col in ("A", "B", "C", "D", "E", "F"):
            ws_summary[f"{col}{row_idx}"].font = font

        for detail in trade.executions:
            ws_detail.append(
                [
                    trade.order_no,
                    trade.stock_name,
                    side_label,
                    detail.exec_time,
                    detail.qty,
                    detail.price,
                    detail.amount,
                    detail.market,
                ]
            )

    buy_total, sell_total = _summary_totals(trade_list)
    total_row = len(trade_list) + 3
    ws_summary[f"A{total_row}"] = "매수 총액"
    ws_summary[f"B{total_row}"] = buy_total
    ws_summary[f"A{total_row + 1}"] = "매도 총액"
    ws_summary[f"B{total_row + 1}"] = sell_total

    for ws in (ws_summary, ws_detail):
        for col in ws.columns:
            width = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 28)

    for row in ws_summary.iter_rows(min_row=2, max_row=ws_summary.max_row, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "#,##0"

    for row in ws_detail.iter_rows(min_row=2, max_row=ws_detail.max_row, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = "#,##0"

    date_text = trade_date
    filename = f"매매일지_{date_text}.xlsx"
    path = output_dir / filename
    wb.save(path)
    return path
